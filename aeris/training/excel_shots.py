"""Make the count file shown in the video, and picture Excel files (count file, stock card) via Excel itself.

python excel_shots.py count            -> comptage_video.xlsx + shots/x_count.png
python excel_shots.py card <xlsx>      -> shots/x_card.png (first stock card sheet)
"""
import os
import sys
import time
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "shots")


def make_count_file(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Product code", "Product", "Lot", "Expiry", "Bin", "Stock AERIS", "Physical QOH", "Comments"])
    rows = [
        ("ACDC-MED-005", "Ringer's lactate 500 ml", "RL0925", date(2027, 9, 30), None, 107, 107, None),
        ("ACDC-MED-005", "Ringer's lactate 500 ml", "RL1130", date(2027, 11, 30), None, 110, 108, "2 flacons cassés"),
        ("ACDC-MED-005", "Ringer's lactate 500 ml", "RL1026", date(2028, 6, 15), None, 60, 60, None),
        ("ACDC-MED-005", "Ringer's lactate 500 ml", "6D340213", date(2029, 3, 1), None, 14700, 14700, None),
    ]
    for r in rows:
        ws.append(list(r))
    head = PatternFill("solid", fgColor="911936")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in range(2, 6):
        ws.cell(r, 4).number_format = "DD/MM/YYYY"
        ws.cell(r, 7).fill = PatternFill("solid", fgColor="FFF4D6")
        ws.cell(r, 7).font = Font(bold=True)
    for col, w in zip("ABCDEFGH", (14, 24, 11, 12, 6, 12, 14, 20)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def picture(xlsx, sheet_index, cell_range, out_png, zoom_to=1700):
    """Excel saves the sheet as HTML (works with Excel hidden); headless Chrome screenshots it."""
    import shutil
    import win32com.client as win32
    from playwright.sync_api import sync_playwright
    out_dir = os.path.join(HERE, "xl_html")
    shutil.rmtree(out_dir, ignore_errors=True)
    os.makedirs(out_dir)
    html = os.path.join(out_dir, "sheet.htm")
    xl = win32.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(os.path.abspath(xlsx), ReadOnly=True)
        ws = wb.Worksheets(sheet_index)
        ws.Activate()
        rng = ws.Range(cell_range) if cell_range else ws.UsedRange
        # xlSourceRange = 4, xlHtmlStatic = 0
        wb.PublishObjects.Add(4, html, ws.Name, rng.Address, 0).Publish(True)
        wb.Close(False)
    finally:
        xl.Quit()
    with sync_playwright() as p:
        b = p.chromium.launch(channel="chrome", headless=True)
        pg = b.new_page(viewport={"width": 1280, "height": 720}, device_scale_factor=1.5)
        pg.goto("file:///" + html.replace(os.sep, "/"))
        pg.add_style_tag(content="body{margin:0;background:#f0f1f4;display:flex;justify-content:center;align-items:center;min-height:620px}"
                                 "table{background:white;box-shadow:0 2px 12px rgba(0,0,0,.15)}")
        pg.wait_for_timeout(300)
        # fit the sheet in the frame (above the caption bar)
        pg.evaluate("""() => { const t = document.querySelector('table');
            const z = Math.min(1.35, 1220 / t.scrollWidth, 560 / t.scrollHeight); t.style.zoom = z; }""")
        pg.wait_for_timeout(500)
        pg.screenshot(path=out_png)
        b.close()
    print("saved", out_png)


if __name__ == "__main__":
    if sys.argv[1] == "count":
        path = os.path.join(HERE, "comptage_video.xlsx")
        make_count_file(path)
        picture(path, 1, "A1:H5", os.path.join(SHOTS, "x_count.png"), zoom_to=1800)
    else:
        picture(sys.argv[2], 2, None, os.path.join(SHOTS, "x_card.png"), zoom_to=1800)
